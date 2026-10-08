# -*- coding: utf-8 -*-
"""Round-trip a page's visible text through an xlsx.

  export  <page.html> <out.xlsx>   pull every editable string out
  apply   <page.html> <in.xlsx>    write the edited strings back

Splits the file into tags and text and keys each text run by index, so the
markup is never touched — only the words between the tags.
"""
import sys, re, io, os

SKIP = re.compile(r'<(script|style)[^>]*>.*?</\1>', re.S)

def tokenize(html):
    """-> list of (is_text, string); joining them reproduces the file exactly."""
    out, pos = [], 0
    # mask script/style so their contents are never offered for editing
    masked = []
    for m in SKIP.finditer(html):
        masked.append((m.start(), m.end()))
    def in_masked(i):
        return any(a <= i < b for a, b in masked)
    for m in re.finditer(r'<[^>]*>', html):
        if m.start() > pos:
            out.append((True, html[pos:m.start()]))
        out.append((False, m.group(0)))
        pos = m.end()
    if pos < len(html):
        out.append((True, html[pos:]))
    # re-mark text inside script/style as non-editable
    idx = 0
    fixed = []
    for is_text, s in out:
        if is_text and in_masked(idx):
            fixed.append((False, s))
        else:
            fixed.append((is_text, s))
        idx += len(s)
    return fixed

def editable(tokens):
    """indices of text runs worth showing a human"""
    keep = []
    for i, (is_text, s) in enumerate(tokens):
        if not is_text: continue
        if not s.strip(): continue
        if len(s.strip()) < 2: continue
        keep.append(i)
    return keep

def section_of(tokens, i):
    """nearest preceding HTML comment banner, as a section label"""
    for j in range(i, -1, -1):
        is_text, s = tokens[j]
        if not is_text and s.startswith('<!--'):
            lab = re.sub(r'<!--\s*=*\s*|\s*=*\s*-->', '', s).strip()
            if lab and 'EQ:' not in lab: return lab
    return ''

def export(page, out_xlsx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    html = io.open(page, encoding='utf-8').read()
    toks = tokenize(html)
    idxs = editable(toks)

    wb = Workbook(); s = wb.active; s.title = "文字"
    hdr = ["#", "區塊", "目前文字", "修改後（在這裡改）", "字數", "備註"]
    widths = [6, 22, 68, 68, 8, 34]
    fill = PatternFill("solid", fgColor="1F3864")
    edit = PatternFill("solid", fgColor="FFFF99")
    thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin,right=thin,top=thin,bottom=thin)
    for c,(h,w) in enumerate(zip(hdr,widths), start=1):
        cell = s.cell(row=1, column=c, value=h)
        cell.font = Font("Arial", 10, bold=True, color="FFFFFF"); cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = bd
        s.column_dimensions[chr(64+c)].width = w
    s.freeze_panes = "A2"; s.row_dimensions[1].height = 28

    r = 2
    for i in idxs:
        text = toks[i][1].strip()
        vals = [i, section_of(toks, i), text, "", len(text), ""]
        for c, v in enumerate(vals, start=1):
            cell = s.cell(row=r, column=c, value=v)
            cell.font = Font("Arial", 10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = bd
            if c == 4: cell.fill = edit
        s.row_dimensions[r].height = max(18, min(90, 15 * (len(text)//34 + 1)))
        r += 1
    wb.save(out_xlsx)
    print(f"✓ {out_xlsx}  —  {len(idxs)} 段文字")

def apply(page, in_xlsx):
    import pandas as pd
    html = io.open(page, encoding='utf-8').read()
    toks = tokenize(html)
    df = pd.read_excel(in_xlsx, sheet_name="文字")
    df.columns = ['i','sec','now','new','len','note'][:len(df.columns)]
    n = 0
    for _, row in df.iterrows():
        new = row['new']
        if not isinstance(new, str) or not new.strip(): continue
        i = int(row['i'])
        old_raw = toks[i][1]
        lead = old_raw[:len(old_raw) - len(old_raw.lstrip())]
        tail = old_raw[len(old_raw.rstrip()):]
        toks[i] = (True, lead + new.strip() + tail)
        n += 1
    io.open(page, 'w', encoding='utf-8').write(''.join(s for _, s in toks))
    print(f"✓ {page}  —  套用 {n} 處修改")

if __name__ == '__main__':
    cmd, page, xl = sys.argv[1], sys.argv[2], sys.argv[3]
    (export if cmd == 'export' else apply)(page, xl)
